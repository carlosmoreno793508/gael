#!/usr/bin/env python3
"""
Genera el análisis de entrenadores + opciones de desarrollo olímpico en:
  - documentos_generados/Analisis_Coaches_Olimpico_Gael.docx  (Word, narrativo)
  - documentos_generados/Coaches_Universidades_Gael.xlsx        (Excel, tabla)

Requiere python-docx y openpyxl.
Uso: python3 scripts/generar_analisis_docx.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from config import DATA, RAIZ  # noqa: E402

SALIDA = RAIZ / "documentos_generados"

# Coaches con pedigrí olímpico de alta confianza (nombre en registro, verificar vigencia)
OLIMPICOS = [
    ("California (Berkeley)", "David Durden", "Head men's EE.UU. Tokio 2020; desarrolló a Ryan Murphy (espalda, oro/récord mundial) y Nathan Adrian (velocidad).", "El mejor caso para un espaldista (Murphy es el espejo)."),
    ("Texas", "Bob Bowman (corrige 'Eddie Reese')", "Entrenó a Michael Phelps y a Léon Marchand; coach olímpico EE.UU.", "Máximo desarrollador vivo. Reach muy alto."),
    ("Florida", "Anthony Nesty", "Oro olímpico (1988); head men's EE.UU. Paris 2024; desarrolló a Caeleb Dressel y Bobby Finke.", "Potencia de velocidad/mariposa; encaja con 50-100 libre y 100 fly."),
    ("Virginia (UVA)", "Todd DeSorbo", "Fábrica de velocistas olímpicos (Kate Douglass, hermanas Walsh); staff olímpico EE.UU.", "Especialista en velocidad: su fortaleza directa."),
    ("Virginia Tech", "Sergio López Miró (VERIFICAR vigencia)", "Medallista olímpico (España 1988); entrenó a Joseph Schooling a oro olímpico 2016.", "Si sigue en VT, top opción alcanzable."),
    ("Stanford / USC / Arizona / Notre Dame / Wisconsin", "Schemmel / Busch / Lindauer / Suguiyama (confianza media)", "Programas con historial olímpico; confirmar nombres actuales.", "Reach; pedigrí de programa."),
]

TIER_A = [  # coaches que NO puedo verificar — solo dato de registro
    "Scott Stern (Cincinnati)", "Pablo Marmolejo (Delaware)", "Iván Sánchez (SIU)",
    "Jim Bossert (Cal Baptist)", "Jack Leavitt (Georgetown)", "Kerry Smith (La Salle)",
    "James Sica (Boston U)", "Mary Ellen Wydan (Lehigh)", "Kristy Jones (Holy Cross)",
    "Vic Riggs (West Virginia)", "Douglas Humphrey (South Dakota State)",
]


def _tier(u: dict) -> str:
    div, conf = u.get("division"), (u.get("conferencia") or "")
    if div == "NCAA D1" and conf == "Ivy":
        return "B — Ivy (sin beca atlética)"
    if div == "NCAA D1" and u.get("prioridad") == "alta":
        return "A — D1 con beca (compite)"
    if div == "NCAA D1":
        return "C — D1 élite (reach)"
    if div == "NCAA D2":
        return "D — D2 (beca, domina)"
    if div == "NAIA":
        return "D — NAIA (beca)"
    return "E — D3 (sin beca)"


def build_docx() -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    h = doc.add_heading("Análisis de entrenadores y desarrollo olímpico — Gael Moreno Sarmiento", level=0)
    sub = doc.add_paragraph("Objetivo: opciones para acercarlo a tiempos de nivel olímpico (LA 2028). "
                            "Sus armas: espalda (100/200) y velocidad libre (50/100) + 100 mariposa.")
    sub.runs[0].italic = True

    aviso = doc.add_paragraph()
    r = aviso.add_run("AVISO DE FIABILIDAD: ")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    aviso.add_run("sin acceso a web en el entorno de generación y con la lista de coaches "
                  "proveniente de un archivo ~2024 (con errores ya detectados, p.ej. Texas). "
                  "Solo se afirman trayectorias de alta confianza; el resto debe verificarse en "
                  "el sitio oficial de cada equipo antes de decidir o contactar.")

    doc.add_paragraph("Nota clave: 'mejor para beca + puntuar en NCAA' NO es lo mismo que 'mejor "
                      "para desarrollo olímpico'. Lo segundo depende del historial del coach "
                      "produciendo internacionales en las pruebas de Gael.")

    doc.add_heading("1. Coaches con pedigrí olímpico real (alta confianza)", level=1)
    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = "Light Grid Accent 1"
    for c, t in zip(tabla.rows[0].cells, ["Programa", "Coach (verificar)", "Trayectoria", "Encaje con Gael"]):
        c.paragraphs[0].add_run(t).bold = True
    for prog, coach, tray, fit in OLIMPICOS:
        cells = tabla.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = prog, coach, tray, fit
    doc.add_paragraph("David Marsh (aparece como associate en Cal en la hoja de Contactos): uno de "
                      "los mejores coaches de velocidad de la historia. Si está en Cal, refuerza a "
                      "Cal para velocidad/espalda. Verificar rol actual.")

    doc.add_heading("2. Coaches Tier A / D2 / D3 — sin bio verificable", level=1)
    doc.add_paragraph("No tengo información fiable de la trayectoria de estos coaches y no la invento. "
                      "Investigar caso por caso (página oficial → Staff; SwimSwam; historial del "
                      "equipo en SwimCloud). Pregunta reveladora en la llamada: «¿A qué nadadores de "
                      "perfil internacional ha desarrollado y con qué progresión de tiempos?»")
    for c in TIER_A:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("3. Mejores opciones para 'tiempos olímpicos 2028'", level=1)
    for i, (t, d) in enumerate([
        ("California (Durden)", "El mejor entorno para un espaldista (Ryan Murphy). Reach alto: seguir bajando tiempos."),
        ("Virginia (DeSorbo) y Florida (Nesty)", "Velocidad de clase mundial; encajan con 50/100 libre y 100 mariposa."),
        ("Virginia Tech (Sergio López, si sigue)", "Mejor equilibrio pedigrí olímpico / accesibilidad. Prioridad de verificación."),
        ("Texas (Bowman) / Stanford / Arizona", "Máximo pedigrí, máximo reach."),
    ], 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{t}: ").bold = True
        p.add_run(d)

    doc.add_heading("Estrategia de dos carriles", level=2)
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Carril desarrollo: ").bold = True
    p.add_run("carta 100% enfocada en espalda + video + progresión de tiempos a 2-3 programas élite (reach).")
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Carril seguro: ").bold = True
    p.add_run("beca en Tier A (Delaware, SIU, Cal Baptist, Patriot) con un coach que confirmes que "
              "desarrolla espaldistas. La relación coach-atleta pesa más que el nombre del programa.")

    doc.add_heading("Matiz para un nadador mexicano", level=2)
    doc.add_paragraph("La vía olímpica pasa por el estándar de clasificación de México y su federación, "
                      "no solo por el sistema NCAA. Elegir un coach que apoye su calendario internacional "
                      "(Mundial Jr., selectivos MEX), no solo el calendario NCAA.")

    doc.add_heading("Verificar de inmediato (cambia el análisis)", level=2)
    for t in ["¿Sergio López sigue en Virginia Tech? (sería la mejor opción alcanzable)",
              "¿David Marsh está en Cal? (refuerza Cal para velocidad/espalda)",
              "Confirmar coach y email vigentes de los Tier A antes de enviar cartas"]:
        doc.add_paragraph(t, style="List Bullet")

    foot = doc.add_paragraph("Documento generado a partir de los datos del proyecto (registro ~2024) + "
                             "conocimiento general de natación NCAA. No sustituye la verificación oficial.")
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].italic = True

    SALIDA.mkdir(exist_ok=True)
    destino = SALIDA / "Analisis_Coaches_Olimpico_Gael.docx"
    doc.save(str(destino))
    return destino


def build_xlsx() -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    universidades = json.loads((DATA / "universidades_objetivo.json").read_text())
    wb = Workbook()
    ws = wb.active
    ws.title = "Coaches"
    cols = ["Universidad", "Div", "Conferencia", "Rank equipo", "Tier", "Coach (registro)",
            "Email", "Teléfono", "Fit tiempos", "Prioridad", "Pedigrí olímpico"]
    ws.append(cols)
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")

    olimp_progs = {"University of California - Berkeley", "Texas", "University of Texas - Austin",
                   "University of Florida", "University of Virginia", "Virginia Tech",
                   "Stanford University", "University of Southern California", "University of Arizona",
                   "University of Notre Dame", "University of Wisconsin"}
    for u in universidades:
        ped = "Sí (verificar)" if u["universidad"] in olimp_progs else ""
        ws.append([u["universidad"], u["division"], u.get("conferencia") or "-",
                   u.get("rank_div") or "-", _tier(u), u.get("coach") or "-",
                   u.get("email") or "-", u.get("telefono") or "-",
                   u.get("fit_tiempos") or "-", u.get("prioridad") or "-", ped])

    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(w, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    SALIDA.mkdir(exist_ok=True)
    destino = SALIDA / "Coaches_Universidades_Gael.xlsx"
    wb.save(str(destino))
    return destino


def main() -> None:
    print("Word:", build_docx())
    print("Excel:", build_xlsx())


if __name__ == "__main__":
    main()
