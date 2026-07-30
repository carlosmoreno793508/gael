#!/usr/bin/env python3
"""
Genera un reporte WHOOP de Gael (recuperación + sueño + strain) en Word,
listo para compartir con coaches, y un resumen en registros/whoop_reporte.md.

Uso:
    # Con datos de ejemplo:
    python3 scripts/whoop_reporte.py

    # Con datos reales de la API (requiere .env con credenciales WHOOP):
    python3 scripts/whoop_reporte.py --live

    # Con un archivo propio (mismas claves: recuperacion/sueno/strain):
    python3 scripts/whoop_reporte.py --datos data/mi_whoop.json
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from statistics import mean

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from config import DATA, REGISTROS, RAIZ, cargar_env   # noqa: E402

SALIDA = RAIZ / "documentos_generados"


def _prom(items, clave):
    vals = [i[clave] for i in items if i.get(clave) is not None]
    return round(mean(vals), 1) if vals else None


def _cargar_datos(args) -> dict:
    if args.live:
        cargar_env()
        from whoop.client import obtener_recuperacion, obtener_sueno, obtener_ciclos
        return {
            "nadador": "Gael Moreno Sarmiento",
            "recuperacion": [r.to_dict() for r in obtener_recuperacion(14)],
            "sueno": [s.to_dict() for s in obtener_sueno(14)],
            "strain": [c.to_dict() for c in obtener_ciclos(14)],
        }
    ruta = Path(args.datos) if args.datos else DATA / "whoop_ejemplo_completo.json"
    return json.loads(ruta.read_text())


def _semaforo(score) -> str:
    if score is None:
        return "-"
    return "Verde" if score >= 67 else "Amarillo" if score >= 34 else "Rojo"


def build_docx(datos: dict) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor

    rec, sue, stn = datos.get("recuperacion", []), datos.get("sueno", []), datos.get("strain", [])
    doc = Document()
    doc.add_heading(f"Reporte WHOOP — {datos.get('nadador', 'Gael Moreno Sarmiento')}", level=0)

    adv = datos.get("ADVERTENCIA")
    if adv:
        p = doc.add_paragraph()
        r = p.add_run("AVISO: " + adv)
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # Resumen
    doc.add_heading("Resumen (promedios del periodo)", level=1)
    resumen = [
        ("Recuperación promedio", f"{_prom(rec, 'recovery_score')} / 100"),
        ("HRV promedio", f"{_prom(rec, 'hrv_ms')} ms"),
        ("FC en reposo promedio", f"{_prom(rec, 'rhr_bpm')} bpm"),
        ("Sueño promedio", f"{_prom(sue, 'horas')} h  ({_prom(sue, 'performance_pct')}% performance)"),
        ("Strain (carga) promedio", f"{_prom(stn, 'strain')} / 21"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light List Accent 1"
    for k, v in resumen:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run(k).bold = True
        c[1].text = str(v)

    # Recuperación
    doc.add_heading("Recuperación diaria (HRV, FC reposo)", level=1)
    tr = doc.add_table(rows=1, cols=5); tr.style = "Light Grid Accent 1"
    for c, h in zip(tr.rows[0].cells, ["Fecha", "Recovery", "Estado", "HRV (ms)", "FC reposo"]):
        c.paragraphs[0].add_run(h).bold = True
    for r in rec:
        cs = tr.add_row().cells
        cs[0].text = r.get("fecha", "-")
        cs[1].text = str(r.get("recovery_score", "-"))
        cs[2].text = _semaforo(r.get("recovery_score"))
        cs[3].text = str(r.get("hrv_ms", "-"))
        cs[4].text = str(r.get("rhr_bpm", "-"))

    # Sueño
    doc.add_heading("Sueño", level=1)
    ts = doc.add_table(rows=1, cols=4); ts.style = "Light Grid Accent 1"
    for c, h in zip(ts.rows[0].cells, ["Fecha", "Horas", "Performance %", "Eficiencia %"]):
        c.paragraphs[0].add_run(h).bold = True
    for s in sue:
        cs = ts.add_row().cells
        cs[0].text = s.get("fecha", "-"); cs[1].text = str(s.get("horas", "-"))
        cs[2].text = str(s.get("performance_pct", "-")); cs[3].text = str(s.get("eficiencia_pct", "-"))

    # Strain
    doc.add_heading("Carga de entrenamiento (Strain)", level=1)
    tc = doc.add_table(rows=1, cols=4); tc.style = "Light Grid Accent 1"
    for c, h in zip(tc.rows[0].cells, ["Fecha", "Strain (0-21)", "FC promedio", "FC máxima"]):
        c.paragraphs[0].add_run(h).bold = True
    for c0 in stn:
        cs = tc.add_row().cells
        cs[0].text = c0.get("fecha", "-"); cs[1].text = str(c0.get("strain", "-"))
        cs[2].text = str(c0.get("fc_promedio", "-")); cs[3].text = str(c0.get("fc_maxima", "-"))

    doc.add_heading("Cómo leerlo (para el coach)", level=1)
    for linea in [
        "Recovery (verde ≥67): listo para carga alta; amarillo: moderar; rojo: priorizar recuperación.",
        "HRV alta y FC en reposo baja = buena adaptación y descanso.",
        "Sueño ≥85% performance sostiene la adaptación al entrenamiento.",
        "Strain diario alto sin recovery suficiente = riesgo de sobrecarga.",
    ]:
        doc.add_paragraph(linea, style="List Bullet")

    foot = doc.add_paragraph("Datos de WHOOP. Para el reporte con métricas reales, "
                             "correr scripts/whoop_reporte.py --live con credenciales en .env.")
    foot.runs[0].font.size = Pt(8); foot.runs[0].italic = True

    SALIDA.mkdir(exist_ok=True)
    destino = SALIDA / "Reporte_WHOOP_Gael.docx"
    doc.save(str(destino))
    return destino


def build_md(datos: dict) -> Path:
    rec, sue, stn = datos.get("recuperacion", []), datos.get("sueno", []), datos.get("strain", [])
    L = [f"# Reporte WHOOP — {datos.get('nadador', 'Gael')}", ""]
    if datos.get("ADVERTENCIA"):
        L += [f"> ⚠️ {datos['ADVERTENCIA']}", ""]
    L += ["## Promedios",
          f"- Recuperación: {_prom(rec,'recovery_score')}/100 | HRV {_prom(rec,'hrv_ms')} ms | "
          f"FC reposo {_prom(rec,'rhr_bpm')} bpm",
          f"- Sueño: {_prom(sue,'horas')} h ({_prom(sue,'performance_pct')}%) | "
          f"Strain: {_prom(stn,'strain')}/21", "",
          "## Recuperación", "| Fecha | Recovery | HRV | FC reposo |", "|---|---|---|---|"]
    for r in rec:
        L.append(f"| {r.get('fecha')} | {r.get('recovery_score')} | {r.get('hrv_ms')} | {r.get('rhr_bpm')} |")
    (REGISTROS / "whoop_reporte.md").write_text("\n".join(L) + "\n")
    return REGISTROS / "whoop_reporte.md"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--live", action="store_true", help="Usar la API real de WHOOP (.env)")
    ap.add_argument("--datos", help="Archivo JSON con recuperacion/sueno/strain")
    args = ap.parse_args()

    datos = _cargar_datos(args)
    md = build_md(datos)
    try:
        docx = build_docx(datos)
        print("Word:", docx)
    except ImportError:
        print("SKIP Word: falta python-docx (pip install python-docx)")
    print("Markdown:", md)
    if datos.get("ADVERTENCIA"):
        print("⚠️ ", datos["ADVERTENCIA"])


if __name__ == "__main__":
    main()
