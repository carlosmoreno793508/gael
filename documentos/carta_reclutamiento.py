"""
Genera una carta de reclutamiento personalizada en Word (.docx) dirigida a un
coach, combinando los datos del atleta y sus mejores tiempos.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class DatosAtleta:
    nombre: str = "Gael Moreno Sarmiento"
    graduacion: str = "2027"
    nacionalidad: str = "México"
    club: str = "Natación Tamaulipas"
    email: str = ""
    telefono: str = ""
    gpa: str = ""
    mejores_tiempos: list[str] = field(default_factory=list)  # "100 Free LCM — 52.34"


def generar_carta(coach: str, universidad: str, atleta: DatosAtleta,
                  destino: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_paragraph(f"Estimado/a Coach {coach},")
    doc.add_paragraph(
        f"Mi nombre es {atleta.nombre}, nadador de {atleta.nacionalidad} "
        f"(club {atleta.club}), generación {atleta.graduacion}. Le escribo por mi "
        f"gran interés en el programa de natación de {universidad} y la posibilidad "
        f"de contribuir a su equipo como estudiante-atleta."
    )
    if atleta.mejores_tiempos:
        doc.add_paragraph("Mis mejores tiempos actuales:")
        for t in atleta.mejores_tiempos:
            doc.add_paragraph(t, style="List Bullet")
    if atleta.gpa:
        doc.add_paragraph(f"Promedio académico (GPA): {atleta.gpa}.")
    doc.add_paragraph(
        "Adjunto mi perfil de SwimCloud y quedo atento a la información sobre su "
        "cuestionario de reclutamiento, estándares de tiempo y opciones de beca."
    )
    doc.add_paragraph("Agradezco su tiempo y consideración.")
    cierre = doc.add_paragraph("\nAtentamente,\n")
    cierre.add_run(atleta.nombre).bold = True
    contacto = " | ".join(x for x in (atleta.email, atleta.telefono) if x)
    if contacto:
        p = doc.add_paragraph(contacto)
        p.runs[0].font.size = Pt(9)
    doc.save(str(destino))
    return destino
