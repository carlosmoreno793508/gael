"""
Genera la tabla de universidades objetivo en Word (.docx) y Excel (.xlsx).

Fuente: data/universidades_objetivo.json (o la salida de correos/analisis.py).
Importaciones perezosas: no requiere python-docx/openpyxl para importar el módulo.
"""
from __future__ import annotations

from pathlib import Path

COLUMNAS = [
    ("universidad", "Universidad"),
    ("division", "División"),
    ("conferencia", "Conferencia"),
    ("estado", "Estado"),
    ("coach", "Coach"),
    ("estatus", "Estatus"),
    ("prioridad", "Prioridad"),
]


def generar_docx(universidades: list[dict], destino: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Universidades objetivo — Gael Moreno Sarmiento", level=1)
    tabla = doc.add_table(rows=1, cols=len(COLUMNAS))
    tabla.style = "Light Grid Accent 1"
    for celda, (_, titulo) in zip(tabla.rows[0].cells, COLUMNAS):
        celda.paragraphs[0].add_run(titulo).bold = True
    for u in universidades:
        fila = tabla.add_row().cells
        for celda, (clave, _) in zip(fila, COLUMNAS):
            celda.text = str(u.get(clave) or "-")
    doc.add_paragraph()
    p = doc.add_paragraph(f"Total: {len(universidades)} universidades.")
    p.runs[0].font.size = Pt(9)
    doc.save(str(destino))
    return destino


def generar_xlsx(universidades: list[dict], destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Universidades"
    ws.append([titulo for _, titulo in COLUMNAS])
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for u in universidades:
        ws.append([u.get(clave) if u.get(clave) is not None else "-"
                   for clave, _ in COLUMNAS])
    for col in ws.columns:
        ancho = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(ancho, 40)
    wb.save(str(destino))
    return destino
